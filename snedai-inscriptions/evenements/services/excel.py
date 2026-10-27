"""
L'export EXCEL : la liste des inscriptions d'un événement, avec la colonne « Présent » (Vrai ou Faux).
On utilise la bibliothèque openpyxl, qui fabrique de vrais fichiers .xlsx.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .formatting import timestamp

# Les colonnes : (titre, largeur).
COLONNES = [('Nom', 22), ('Prénom', 22), ('Filiale', 24), ('Numéro de téléphone', 22), ('Adresse e-mail', 34),
            ('Présent', 12), ('Inscrit le', 20), ('Scanné le', 20)]


def build_inscriptions_workbook(event, inscriptions):
    """Fabrique le fichier Excel et le renvoie sous forme d'octets (prêt à être téléchargé)."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = 'Inscriptions'
    feuille.append([titre for titre, _ in COLONNES])
    for i in inscriptions:
        feuille.append([i.nom, i.prenom, i.filiale, i.telephone, i.email,
                        'Vrai' if i.present else 'Faux',               # la colonne vrai / faux demandée
                        timestamp(i.created_at), timestamp(i.scanne_at)])
    for numero, (_, largeur) in enumerate(COLONNES):
        feuille.column_dimensions[chr(65 + numero)].width = largeur       # A, B, C…
    for cellule in feuille[1]:                                             # en-tête en bleu
        cellule.font = Font(bold=True, color='FFFFFFFF')
        cellule.fill = PatternFill('solid', fgColor='FF0B3D91')
    feuille.freeze_panes = 'A2'                                            # la première ligne reste visible
    if inscriptions:
        feuille.auto_filter.ref = feuille.dimensions                       # filtres automatiques

    # Une seconde feuille avec les informations de l'événement.
    resume = classeur.create_sheet('Événement')
    for ligne in (['Événement', event.titre], ['Date', event.range_text], ['Lieu', event.lieu or ''],
                  ['Inscrits', len(inscriptions)], ['Présents (QR scanné)', sum(1 for i in inscriptions if i.present)]):
        resume.append(ligne)
    resume.column_dimensions['A'].width = 26
    resume.column_dimensions['B'].width = 60

    tampon = io.BytesIO()                                                  # un « fichier » en mémoire
    classeur.save(tampon)
    return tampon.getvalue()
