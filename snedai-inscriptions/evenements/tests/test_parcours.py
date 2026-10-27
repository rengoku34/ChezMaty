"""Tests automatiques : on simule un navigateur pour vérifier tout le parcours demandé par l'encadreur."""
import io
import json
import time
from unittest import mock

from django.core.cache import cache
from django.test import Client, TestCase, override_settings
from openpyxl import load_workbook

from ..models import Event, Inscription, User
from ..services.mailer import reset_mailer

REGLAGES = dict(ADMIN_EMAILS=['admin@test.ci'], FILIALES=['SNEDAI', 'SNEDAI Telecom'], AGENTMAIL_API_KEY='cle',
                AGENTMAIL_INBOX='snedai@agentmail.to', BREVO_API_KEY='', SMTP_HOST='', LOGIN_RATE_LIMIT=1000,
                PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
FORMULAIRE = {'nom': 'Kone', 'prenom': 'Awa', 'filiale': 'SNEDAI Telecom', 'telephone': '07 07 11 22 33', 'email': 'Awa.Kone@exemple.ci'}
MDP = 'motdepasse1'


class Reponse:
    ok, status_code, text = True, 200, ''


@override_settings(**REGLAGES)
class BaseTest(TestCase):
    def setUp(self):
        cache.clear()
        reset_mailer()
        self.envoyes, self.panne = [], False

        def faux_envoi(url, json=None, **kw):
            if self.panne:
                r = Reponse(); r.ok, r.status_code, r.text = False, 401, 'Clé refusée'
                return r
            self.envoyes.append(json)
            return Reponse()

        patch = mock.patch('evenements.services.mailer.requests.post', side_effect=faux_envoi)
        patch.start()
        self.addCleanup(patch.stop)
        self.addCleanup(reset_mailer)
        self.admin = User.objects.create_user('admin@test.ci', MDP, nom='Nabaloum', prenom='Aboubakar')
        self.agent = User.objects.create_user('moussa@test.ci', MDP, nom='Traore', prenom='Moussa', is_agent=True)
        self.event = Event.objects.create(titre='Journée bien-être', description='Texte', debut='2099-11-14T09:00',
                                          fin='2099-11-14T16:00', lieu='Siège')

    def connecter(self, email='admin@test.ci'):
        client = Client()
        client.post('/connexion', {'email': email, 'password': MDP})
        return client

    def attendre_emails(self, nombre):
        for _ in range(60):          # l'envoi se fait en arrière-plan
            if len(self.envoyes) >= nombre:
                break
            time.sleep(0.05)


class ParcoursTests(BaseTest):
    def test_parcours_complet_inscription_qr_scan_excel(self):
        # 1. Un employé sans compte ouvre le lien et remplit le formulaire
        publique = Client()
        page = publique.get(f'/e/{self.event.lien}')
        self.assertContains(page, 'Journée bien-être')
        reponse = publique.post(f'/e/{self.event.lien}', FORMULAIRE)
        inscription = Inscription.objects.get()
        self.assertEqual(reponse['Location'], f'/billet/{inscription.qr_token}?nouveau=1')
        self.assertEqual((inscription.present, inscription.email), (False, 'awa.kone@exemple.ci'))   # FAUX au départ

        # 2. Il reçoit un e-mail avec son QR code en pièce jointe
        self.attendre_emails(1)
        self.assertEqual(self.envoyes[0]['to'], ['awa.kone@exemple.ci'])
        self.assertEqual(self.envoyes[0]['attachments'][0]['filename'], 'qrcode.png')
        self.assertIn(inscription.url_billet, self.envoyes[0]['html'])

        # 3. La page du billet affiche le QR (seulement avec le code secret)
        self.assertContains(publique.get(f'/billet/{inscription.qr_token}'), 'data:image/png;base64,')
        self.assertEqual(publique.get('/billet/inconnu').status_code, 404)
        self.assertEqual(publique.get(f'/billet/{inscription.qr_token}/qr.png')['Content-Type'], 'image/png')

        # 4. Un agent scanne : « présent » passe à VRAI, une seule fois
        agent = self.connecter('moussa@test.ci')
        code = f'SNEDAI-EVT:{inscription.qr_token}'
        premier = agent.post('/scan', json.dumps({'code': code}), content_type='application/json')
        self.assertEqual((premier.status_code, premier.json()['status'], premier.json()['person']['prenom']), (200, 'present', 'Awa'))
        second = agent.post('/scan', json.dumps({'code': code}), content_type='application/json')
        self.assertEqual((second.status_code, second.json()['status']), (409, 'deja'))
        inscription.refresh_from_db()
        self.assertEqual((inscription.present, inscription.scanne_par), (True, self.agent))
        inconnu = agent.post('/scan', json.dumps({'code': 'SNEDAI-EVT:' + 'z' * 32}), content_type='application/json')
        self.assertEqual((inconnu.status_code, inconnu.json()['status']), (404, 'inconnu'))
        self.assertContains(publique.get(f'/billet/{inscription.qr_token}'), 'Présence validée')

        # 5. L'administrateur voit la liste (Vrai) et l'exporte en Excel
        admin = self.connecter()
        detail = admin.get(f'/admin/evenements/{self.event.id}')
        self.assertContains(detail, 'Kone')
        self.assertContains(detail, '>Vrai<')
        classeur = load_workbook(io.BytesIO(admin.get(f'/admin/evenements/{self.event.id}/export.xlsx').content))
        lignes = list(classeur['Inscriptions'].iter_rows(values_only=True))
        self.assertEqual(lignes[0][:6], ('Nom', 'Prénom', 'Filiale', 'Numéro de téléphone', 'Adresse e-mail', 'Présent'))
        self.assertEqual(lignes[1][:6], ('Kone', 'Awa', 'SNEDAI Telecom', '07 07 11 22 33', 'awa.kone@exemple.ci', 'Vrai'))

    def test_validation_du_formulaire(self):
        for changement, message in (({'nom': ''}, 'nom et votre prénom'), ({'filiale': 'Autre'}, 'filiale dans la liste'),
                                    ({'telephone': 'abc'}, 'téléphone invalide'), ({'email': 'pasunemail'}, 'e-mail invalide')):
            reponse = Client().post(f'/e/{self.event.lien}', {**FORMULAIRE, **changement})
            self.assertContains(reponse, message, status_code=400)
        self.assertEqual(Inscription.objects.count(), 0)

    def test_une_seule_inscription_par_adresse_et_pas_de_fuite_du_billet(self):
        client = Client()
        client.post(f'/e/{self.event.lien}', FORMULAIRE)
        reponse = client.post(f'/e/{self.event.lien}', {**FORMULAIRE, 'email': 'AWA.KONE@exemple.ci', 'nom': 'Autre'})
        self.assertEqual((reponse.status_code, Inscription.objects.count()), (200, 1))
        self.assertContains(reponse, 'déjà inscrite')
        self.assertNotContains(reponse, 'data:image/png')          # on ne montre PAS le QR à quelqu'un qui connaît juste l'e-mail
        self.attendre_emails(2)                                    # mais le billet est renvoyé par e-mail
        self.assertEqual(len(self.envoyes), 2)
        # la même adresse peut s'inscrire à un AUTRE événement
        autre = Event.objects.create(titre='Autre', description='x')
        Client().post(f'/e/{autre.lien}', FORMULAIRE)
        self.assertEqual(Inscription.objects.count(), 2)

    def test_une_panne_d_e_mail_n_empeche_pas_l_inscription(self):
        self.panne = True
        reponse = Client().post(f'/e/{self.event.lien}', FORMULAIRE)
        self.assertEqual((reponse.status_code, Inscription.objects.count()), (302, 1))
        time.sleep(0.3)
        admin = self.connecter()
        self.assertContains(admin.get('/admin/agents'), 'Clé refusée')

    def test_lien_inconnu(self):
        self.assertEqual(Client().get('/e/nimportequoi').status_code, 404)


class BackofficeTests(BaseTest):
    def test_crud_des_evenements(self):
        admin = self.connecter()
        # Create
        reponse = admin.post('/admin/evenements/creer', {'titre': 'Soirée', 'description': 'Venez', 'lieu': 'Siège', 'debut': '2099-12-20T18:00', 'fin': '2099-12-20T23:00'})
        nouveau = Event.objects.get(titre='Soirée')
        self.assertEqual(reponse['Location'], f'/admin/evenements/{nouveau.id}')
        self.assertTrue(nouveau.lien)
        # Read
        self.assertContains(admin.get('/admin/evenements'), 'Soirée')
        self.assertContains(admin.get(f'/admin/evenements/{nouveau.id}'), f'/e/{nouveau.lien}')
        # Update : le lien ne change pas
        ancien_lien = nouveau.lien
        admin.post(f'/admin/evenements/{nouveau.id}/modifier', {'titre': 'Soirée 2099', 'description': 'Nouveau', 'lieu': 'Salle'})
        nouveau.refresh_from_db()
        self.assertEqual((nouveau.titre, nouveau.lien), ('Soirée 2099', ancien_lien))
        # Delete : efface aussi les inscriptions
        Client().post(f'/e/{nouveau.lien}', FORMULAIRE)
        self.assertEqual(Inscription.objects.count(), 1)
        admin.post(f'/admin/evenements/{nouveau.id}/supprimer')
        self.assertFalse(Event.objects.filter(pk=nouveau.pk).exists())
        self.assertEqual(Inscription.objects.count(), 0)

    def test_validation_du_formulaire_d_evenement(self):
        admin = self.connecter()
        for champs, message in (({'titre': '', 'description': 'x'}, 'obligatoires'), ({'titre': 'a', 'description': 'b', 'debut': 'demain'}, 'Horaires invalides'),
                                ({'titre': 'a', 'description': 'b', 'fin': '2099-01-01T10:00'}, 'avant l’heure de fin'),
                                ({'titre': 'a', 'description': 'b', 'debut': '2099-01-02T10:00', 'fin': '2099-01-01T10:00'}, 'après le début')):
            self.assertContains(admin.post('/admin/evenements/creer', champs), message, status_code=400)

    def test_droits_d_acces(self):
        # Sans connexion : renvoi à la connexion
        for chemin in ('/admin/evenements', f'/admin/evenements/{self.event.id}', '/admin/agents', '/scan'):
            self.assertEqual((Client().get(chemin).status_code, Client().get(chemin)['Location']), (302, '/connexion'), chemin)
        # Un agent ne voit que le scan
        agent = self.connecter('moussa@test.ci')
        self.assertEqual(agent.get('/scan').status_code, 200)
        for chemin in ('/admin/evenements', f'/admin/evenements/{self.event.id}', '/admin/agents', f'/admin/evenements/{self.event.id}/export.xlsx'):
            self.assertEqual(agent.get(chemin).status_code, 403, chemin)
        self.assertEqual(agent.post('/admin/evenements/creer', {'titre': 'x', 'description': 'y'}).status_code, 403)
        self.assertEqual(agent.post(f'/admin/evenements/{self.event.id}/supprimer').status_code, 403)
        self.assertTrue(Event.objects.filter(pk=self.event.pk).exists())

    def test_connexion(self):
        self.assertEqual(Client().post('/connexion', {'email': 'admin@test.ci', 'password': 'faux12345'}).status_code, 401)
        self.assertEqual(Client().post('/connexion', {'email': 'inconnu@test.ci', 'password': MDP}).status_code, 401)
        self.assertEqual(Client().post('/connexion', {'email': ' ADMIN@test.ci ', 'password': MDP})['Location'], '/admin/evenements')
        self.assertEqual(Client().post('/connexion', {'email': 'moussa@test.ci', 'password': MDP})['Location'], '/scan')
        admin = self.connecter()
        reponse = admin.post('/deconnexion')
        self.assertEqual(reponse['Clear-Site-Data'], '"cache"')
        self.assertEqual(admin.get('/admin/evenements')['Location'], '/connexion')
        self.assertEqual(Client(enforce_csrf_checks=True).post('/connexion', {'email': 'a', 'password': 'b'}).status_code, 403)

    def test_gestion_des_agents(self):
        admin = self.connecter()
        reponse = admin.post('/admin/agents', {'nom': 'Yao', 'prenom': 'Marc', 'email': 'marc@test.ci', 'password': 'motdepasse2'})
        self.assertEqual(reponse['Location'], '/admin/agents')
        nouveau = User.objects.get(email='marc@test.ci')
        self.assertTrue(nouveau.is_agent and not nouveau.is_admin)
        self.assertContains(admin.post('/admin/agents', {'nom': 'Y', 'prenom': 'M', 'email': 'marc@test.ci', 'password': 'motdepasse2'}), 'existe déjà', status_code=400)
        self.assertContains(admin.post('/admin/agents', {'nom': 'Y', 'prenom': 'M', 'email': 'x@test.ci', 'password': 'court'}), '8 caractères', status_code=400)
        admin.post(f'/admin/agents/{self.admin.id}/supprimer')
        self.assertTrue(User.objects.filter(pk=self.admin.pk).exists())          # l'administrateur est protégé
        admin.post(f'/admin/agents/{nouveau.id}/supprimer')
        self.assertFalse(User.objects.filter(pk=nouveau.pk).exists())

    def test_supprimer_une_inscription_et_test_d_e_mail(self):
        client = Client()
        client.post(f'/e/{self.event.lien}', FORMULAIRE)
        admin = self.connecter()
        admin.post(f'/admin/evenements/{self.event.id}/inscriptions/{Inscription.objects.get().id}/supprimer')
        self.assertEqual(Inscription.objects.count(), 0)
        self.attendre_emails(1)
        self.envoyes.clear()
        self.assertContains(admin.post('/admin/test-email', follow=True), 'E-mail de test envoyé')
        self.assertEqual(self.envoyes[0]['to'], ['admin@test.ci'])

    def test_securite_des_pages_publiques(self):
        reponse = Client(enforce_csrf_checks=True).post(f'/e/{self.event.lien}', FORMULAIRE)
        self.assertEqual(reponse.status_code, 403)                  # formulaire sans jeton CSRF refusé
        with override_settings(LOGIN_RATE_LIMIT=2):
            cache.clear()
            client = Client()
            codes = [client.post(f'/e/{self.event.lien}', {**FORMULAIRE, 'email': f'a{i}@exemple.ci'}).status_code for i in range(4)]
            self.assertEqual(codes[-1], 429)                        # trop de tentatives depuis la même adresse IP
