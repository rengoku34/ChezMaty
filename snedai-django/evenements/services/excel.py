"""Export Excel : un résumé et les trois listes du tableau de bord, ou une seule liste."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .formatting import timestamp

PERSON = [('Nom', 22), ('Prénom', 22), ('Adresse e-mail', 34), ('Filiale', 24)]


def _person(user):
    return [user.nom, user.prenom, user.email, user.filiale]


# Les trois listes : nom de la feuille, colonnes, et fonction qui fabrique une ligne.
LISTS = {
    'inscrits': ('Inscrits sur le site', PERSON + [('Inscrit le', 20)], lambda u: _person(u) + [timestamp(u.created_at)]),
    'confirmes': (
        'Présents annoncés',
        PERSON + [('Confirmé le', 20), ('Était là (QR scanné)', 20)],
        lambda p: _person(p.user) + [timestamp(p.confirme_at), 'Oui' if p.scanne_at else 'Non'],
    ),
    'presents': ('Présents confirmés (QR)', PERSON + [('Scanné le', 20)], lambda p: _person(p.user) + [timestamp(p.scanne_at)]),
}
LIST_NAMES = list(LISTS)


def _sheet(workbook, name, columns, rows, filter_rows=True):
    sheet = workbook.create_sheet(name)
    sheet.append([header for header, _ in columns])
    for row in rows:
        sheet.append(row)
    for index, (_, width) in enumerate(columns):
        sheet.column_dimensions[chr(65 + index)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFFFF')
        cell.fill = PatternFill('solid', fgColor='FF0B3D91')
    sheet.freeze_panes = 'A2'
    if rows and filter_rows:
        sheet.auto_filter.ref = sheet.dimensions
    return sheet


def build_event_workbook(event, inscrits, confirmes, presents, only=None):
    workbook = Workbook()
    workbook.remove(workbook.active)
    rows = {'inscrits': inscrits, 'confirmes': confirmes, 'presents': presents}

    if only:
        name, columns, make = LISTS[only]
        _sheet(workbook, name, columns, [make(r) for r in rows[only]])
    else:
        _sheet(
            workbook, 'Résumé', [('Information', 30), ('Valeur', 60)],
            [
                ['Événement', event.titre],
                ['Date', event.range_text],
                ['Lieu', event.lieu or ''],
                ['Statut', f'Terminé ({timestamp(event.termine_at)})' if event.termine_at else 'En cours'],
                ['Inscrits sur le site', len(inscrits)],
                ['Ont dit qu’ils seront présents', len(confirmes)],
                ['Présents confirmés (QR scanné)', len(presents)],
            ],
            filter_rows=False,
        )
        for key, (name, columns, make) in LISTS.items():
            _sheet(workbook, name, columns, [make(r) for r in rows[key]])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
