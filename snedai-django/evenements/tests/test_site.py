import re
from unittest import mock

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, override_settings
from django.utils import timezone

from ..models import Event, Participation, User
from .helpers import PASSWORD, BaseTest

PNG = (b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89'
       b'\x00\x00\x00\rIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02\xfe\xa7\x9a\xa0\xa0\x00\x00\x00\x00IEND\xaeB`\x82')


class AccessTests(BaseTest):
    def test_sans_compte_on_ne_voit_pas_les_evenements(self):
        event = self.event()
        for path in ('/', f'/evenements/{event.id}', f'/evenements/{event.id}/photo', '/scan', '/admin/tableau-de-bord'):
            response = Client().get(path)
            self.assertEqual((response.status_code, response['Location']), (302, '/connexion'), path)

    def test_pages_legales_publiques(self):
        for path in ('/confidentialite', '/conditions'):
            self.assertEqual(Client().get(path).status_code, 200)

    def test_apres_deconnexion_plus_rien_n_est_visible_ni_garde_en_cache(self):
        self.user('awa@test.ci')
        client = Client()
        self.login(client, 'awa@test.ci')
        page = client.get('/')
        self.assertEqual(page.status_code, 200)
        self.assertIn('no-store', page['Cache-Control'])
        response = client.post('/deconnexion')
        self.assertEqual(response['Clear-Site-Data'], '"cache"')
        self.assertEqual(client.get('/')['Location'], '/connexion')

    def test_mauvais_mot_de_passe_refuse(self):
        self.user('awa@test.ci')
        for email, password in (('awa@test.ci', 'faux12345'), ('inconnu@test.ci', PASSWORD)):
            response = Client().post('/connexion', {'email': email, 'password': password})
            self.assertEqual(response.status_code, 401)
            self.assertContains(response, 'incorrect', status_code=401)

    def test_connexion_insensible_a_la_casse_de_l_adresse(self):
        self.user('awa@test.ci')
        self.assertEqual(Client().post('/connexion', {'email': ' AWA@Test.ci ', 'password': PASSWORD})['Location'], '/')

    def test_formulaire_sans_jeton_csrf_refuse(self):
        response = Client(enforce_csrf_checks=True).post('/connexion', {'email': 'a@test.ci', 'password': 'x'})
        self.assertEqual(response.status_code, 403)

    def test_inscription_demande_une_adresse_valide_et_nouvelle(self):
        self.user('awa@test.ci')
        base = {'nom': 'Kone', 'prenom': 'Issa', 'filiale': 'SNEDAI', 'password': PASSWORD, 'password_confirmation': PASSWORD, 'methode': 'email'}
        for changes, message in (
            ({'email': 'pas-une-adresse'}, 'Adresse e-mail invalide'),
            ({'email': 'issa@test.ci', 'password_confirmation': 'autre1234'}, 'pas identiques'),
            ({'email': 'issa@test.ci', 'password': 'court', 'password_confirmation': 'court'}, 'au moins 8 caractères'),
            ({'email': 'AWA@test.ci'}, 'existe déjà'),
        ):
            response = Client().post('/inscription', {**base, **changes})
            self.assertContains(response, message, status_code=400)
        self.assertEqual(User.objects.count(), 1)

    @override_settings(FILIALES=['SNEDAI', 'SNEDAI Telecom'])
    def test_la_filiale_doit_etre_dans_la_liste(self):
        response = Client().post('/inscription', {'nom': 'K', 'prenom': 'I', 'filiale': 'Autre', 'email': 'i@test.ci', 'password': PASSWORD, 'password_confirmation': PASSWORD})
        self.assertContains(response, 'Choisissez une filiale', status_code=400)

    def test_seul_l_administrateur_accede_au_tableau_de_bord_a_la_publication_et_a_l_export(self):
        event = self.event()
        self.user('awa@test.ci')
        self.user('admin@test.ci', nom='Admin')
        employee, admin = Client(), Client()
        self.login(employee, 'awa@test.ci')
        self.login(admin, 'admin@test.ci')
        for path in ('/admin/tableau-de-bord', '/admin/evenements/nouveau', '/admin/agents', f'/admin/evenements/{event.id}/export.xlsx', f'/admin/evenements/{event.id}/rapport.pdf'):
            self.assertEqual(employee.get(path).status_code, 403, path)
            self.assertEqual(admin.get(path).status_code, 200, path)
        self.assertEqual(employee.post('/admin/evenements', {'titre': 'x', 'description': 'y'}).status_code, 403)

    def test_le_menu_n_affiche_que_les_liens_autorises(self):
        self.user('awa@test.ci')
        self.user('moussa@test.ci', agent=True)
        self.user('admin@test.ci')
        expected = {'awa@test.ci': (False, False), 'moussa@test.ci': (True, False), 'admin@test.ci': (True, True)}
        for email, (scan, admin) in expected.items():
            client = Client()
            self.login(client, email)
            html = client.get('/').content.decode()
            self.assertEqual('href="/scan"' in html, scan, email)
            self.assertEqual('href="/admin/tableau-de-bord"' in html, admin, email)

    def test_photos_accessibles_seulement_une_fois_connecte(self):
        event = self.event()
        Event.objects.filter(pk=event.pk).update(photo_data=PNG, photo_type='image/png')
        self.assertEqual(Client().get(f'/evenements/{event.id}/photo').status_code, 302)
        self.user('awa@test.ci')
        client = Client()
        self.login(client, 'awa@test.ci')
        response = client.get(f'/evenements/{event.id}/photo')
        self.assertEqual((response.status_code, response['Content-Type']), (200, 'image/png'))
        self.assertEqual(response.content, PNG)

    def test_application_installable_manifeste_et_service_worker(self):
        manifest = Client().get('/manifest.webmanifest')
        self.assertEqual(manifest['Content-Type'], 'application/manifest+json')
        self.assertEqual(manifest.json()['display'], 'standalone')
        worker = Client().get('/sw.js')
        self.assertIn('javascript', worker['Content-Type'])
        self.assertEqual(worker['Cache-Control'], 'no-cache')
        self.assertEqual(Client().get('/.well-known/assetlinks.json').status_code, 404)
        with override_settings(ANDROID_PACKAGE_NAME='ci.snedai.events', ANDROID_CERT_SHA256=['AA:BB']):
            self.assertEqual(Client().get('/.well-known/assetlinks.json').json()[0]['target']['package_name'], 'ci.snedai.events')

    def test_en_tetes_de_securite(self):
        response = Client().get('/connexion')
        self.assertIn("script-src 'self'", response['Content-Security-Policy'])
        self.assertEqual(response['X-Content-Type-Options'], 'nosniff')

    def test_page_introuvable_et_evenement_inexistant(self):
        self.user('awa@test.ci')
        client = Client()
        self.login(client, 'awa@test.ci')
        self.assertEqual(client.get('/evenements/9999').status_code, 404)
        self.assertEqual(client.get('/nimporte-quoi').status_code, 404)


class FullFlowTests(BaseTest):
    def test_parcours_complet_publication_presence_qr_scan_tableau_de_bord_excel_termine(self):
        self.user('admin@test.ci', nom='Admin', prenom='Boss')
        self.user('awa@test.ci', nom='Kone', prenom='Awa', filiale='SNEDAI Telecom')
        self.user('moussa@test.ci', nom='Traore', prenom='Moussa', agent=True)
        admin, awa, moussa = Client(), Client(), Client()
        for client, email in ((admin, 'admin@test.ci'), (awa, 'awa@test.ci'), (moussa, 'moussa@test.ci')):
            self.login(client, email)

        # Publication avec photo : chaque inscrit est prévenu
        response = admin.post('/admin/evenements', {
            'titre': 'Soirée de fin d’année', 'description': 'Venez nombreux.', 'lieu': 'Siège', 'debut': '2099-12-20T18:00', 'fin': '2099-12-20T23:00',
            'photo': SimpleUploadedFile('photo.png', PNG, content_type='image/png'),
        })
        event = Event.objects.get()
        self.assertEqual(response['Location'], f'/admin/tableau-de-bord?evenement={event.id}')
        self.assertTrue(event.has_photo)
        import time
        for _ in range(50):  # envoi en arrière-plan
            if len(self.sent) >= 3:
                break
            time.sleep(0.05)
        self.assertEqual(sorted(m['to'][0] for m in self.sent), ['admin@test.ci', 'awa@test.ci', 'moussa@test.ci'])
        self.assertIn('Nouvel événement', self.sent[0]['subject'])
        self.sent.clear()

        # L'employé voit l'événement et confirme : QR par e-mail (pièce jointe) et à l'écran
        self.assertContains(awa.get('/'), 'Soirée de fin d’année')
        self.assertContains(awa.get(f'/evenements/{event.id}'), 'Je serai présent(e)')
        awa.post(f'/evenements/{event.id}/participer')
        participation = Participation.objects.get()
        self.assertEqual(len(self.sent), 1)
        self.assertEqual(self.sent[0]['attachments'][0]['filename'], 'qrcode.png')
        page = awa.get(f'/evenements/{event.id}').content.decode()
        self.assertIn('data:image/png;base64,', page)
        self.assertEqual(awa.get(f'/evenements/{event.id}/qr.png')['Content-Type'], 'image/png')
        awa.post(f'/evenements/{event.id}/participer')  # un second clic ne renvoie rien
        self.assertEqual((len(self.sent), Participation.objects.count()), (1, 1))

        # Agenda
        ics = awa.get(f'/evenements/{event.id}/agenda.ics')
        self.assertIn('DTSTART:20991220T180000Z', ics.content.decode())
        self.assertIn('google.com/calendar', awa.get(f'/evenements/{event.id}').content.decode())

        # Scan : seul un agent, une seule fois
        code = f'SNEDAI-EVT:{participation.qr_token}'
        self.assertEqual(awa.post('/scan', {'code': code}, content_type='application/json').status_code, 403)
        first = moussa.post('/scan', {'code': code}, content_type='application/json')
        self.assertEqual(first.status_code, 200)
        self.assertEqual(first.json()['person']['prenom'], 'Awa')
        second = moussa.post('/scan', {'code': code}, content_type='application/json')
        self.assertEqual((second.status_code, second.json()['status']), (409, 'deja'))
        self.assertEqual(moussa.post('/scan', {'code': 'SNEDAI-EVT:' + 'x' * 32}, content_type='application/json').status_code, 404)
        self.assertEqual(awa.get(f'/evenements/{event.id}').content.decode().count('Présence validée'), 1)

        # Tableau de bord, Excel et PDF
        board = admin.get(f'/admin/tableau-de-bord?evenement={event.id}')
        self.assertContains(board, 'Kone')
        self.assertContains(board, 'SNEDAI Telecom')
        self.assertEqual(board.context['presents_count'], 1)
        self.assertEqual(board.context['rates'], {'reponse': 33, 'presence': 100})
        from openpyxl import load_workbook
        import io
        book = load_workbook(io.BytesIO(admin.get(f'/admin/evenements/{event.id}/export.xlsx').content))
        self.assertEqual(book.sheetnames, ['Résumé', 'Inscrits sur le site', 'Présents annoncés', 'Présents confirmés (QR)'])
        self.assertEqual(book['Présents confirmés (QR)']['A2'].value, 'Kone')
        single = load_workbook(io.BytesIO(admin.get(f'/admin/evenements/{event.id}/export.xlsx?liste=presents').content))
        self.assertEqual(single.sheetnames, ['Présents confirmés (QR)'])
        pdf = admin.get(f'/admin/evenements/{event.id}/rapport.pdf')
        self.assertEqual(pdf['Content-Type'], 'application/pdf')
        self.assertTrue(pdf.content.startswith(b'%PDF'))

        # Modification, puis fin : plus de scan, plus de réponses
        admin.post(f'/admin/evenements/{event.id}/modifier', {'titre': 'Soirée 2099', 'description': 'Nouveau texte', 'lieu': 'Salle', 'debut': '2099-12-20T19:00', 'fin': '2099-12-20T23:00'})
        event.refresh_from_db()
        self.assertEqual((event.titre, event.has_photo), ('Soirée 2099', True))
        admin.post(f'/admin/evenements/{event.id}/terminer')
        event.refresh_from_db()
        self.assertIsNotNone(event.termine_at)
        self.assertEqual(moussa.post('/scan', {'code': code}, content_type='application/json').json()['status'], 'termine')
        self.assertContains(awa.get(f'/evenements/{event.id}'), 'Cet événement est terminé')
        self.assertContains(awa.get('/'), 'Événements passés')

    def test_validation_du_formulaire_d_evenement(self):
        self.user('admin@test.ci')
        admin = Client()
        self.login(admin, 'admin@test.ci')
        for fields, message in (
            ({'titre': '', 'description': 'x'}, 'obligatoires'),
            ({'titre': 'a', 'description': 'b', 'debut': 'demain'}, 'Horaires invalides'),
            ({'titre': 'a', 'description': 'b', 'fin': '2099-01-01T10:00'}, 'avant l’heure de fin'),
            ({'titre': 'a', 'description': 'b', 'debut': '2099-01-02T10:00', 'fin': '2099-01-01T10:00'}, 'après le début'),
        ):
            self.assertContains(admin.post('/admin/evenements', fields), message, status_code=400)
        gros = SimpleUploadedFile('gros.png', b'0' * (3 * 1024 * 1024 + 1), content_type='image/png')
        self.assertContains(admin.post('/admin/evenements', {'titre': 'a', 'description': 'b', 'photo': gros}), '3 Mo', status_code=400)
        faux = SimpleUploadedFile('script.html', b'<script>', content_type='text/html')
        self.assertContains(admin.post('/admin/evenements', {'titre': 'a', 'description': 'b', 'photo': faux}), 'Photo invalide', status_code=400)
        self.assertEqual(Event.objects.count(), 0)

    def test_annuler_sa_participation(self):
        event = self.event()
        self.user('awa@test.ci')
        awa = Client()
        self.login(awa, 'awa@test.ci')
        awa.post(f'/evenements/{event.id}/participer')
        self.assertEqual(Participation.objects.count(), 1)
        awa.post(f'/evenements/{event.id}/annuler')
        self.assertEqual(Participation.objects.count(), 0)

    def test_compte_a_rebours_vers_le_prochain_evenement(self):
        self.event('Plus tard', debut='2099-12-01T09:00', fin='2099-12-01T10:00')
        self.event('Bientôt', debut='2099-11-01T09:00', fin='2099-11-01T10:00')
        self.user('awa@test.ci')
        awa = Client()
        self.login(awa, 'awa@test.ci')
        html = awa.get('/').content.decode()
        self.assertIn('data-countdown="2099-11-01T09:00:00.000Z"', html)
        self.assertRegex(html, r'Prochain événement</span>\s*<h2>Bientôt')


class AccountTests(BaseTest):
    def test_seul_l_administrateur_supprime_des_comptes_et_jamais_le_sien(self):
        admin_user = self.user('admin@test.ci')
        awa = self.user('awa@test.ci')
        moussa = self.user('moussa@test.ci')
        admin, employee = Client(), Client()
        self.login(admin, 'admin@test.ci')
        self.login(employee, 'awa@test.ci')
        self.assertEqual(employee.post(f'/admin/comptes/{moussa.id}/supprimer').status_code, 403)
        admin.post(f'/admin/comptes/{admin_user.id}/supprimer')
        self.assertTrue(User.objects.filter(pk=admin_user.pk).exists())
        admin.post(f'/admin/comptes/{awa.id}/supprimer')
        self.assertFalse(User.objects.filter(pk=awa.pk).exists())
        # la personne supprimée est déconnectée
        self.assertEqual(employee.get('/')['Location'], '/connexion')

    def test_nommer_un_agent(self):
        self.user('admin@test.ci')
        awa = self.user('awa@test.ci')
        admin = Client()
        self.login(admin, 'admin@test.ci')
        admin.post(f'/admin/agents/{awa.id}', {'agent': '1'})
        awa.refresh_from_db()
        self.assertTrue(awa.is_agent)
        admin.post(f'/admin/agents/{awa.id}', {'agent': '0'})
        awa.refresh_from_db()
        self.assertFalse(awa.is_agent)

    def test_une_panne_d_envoi_est_memorisee_et_le_test_d_envoi_affiche_l_erreur(self):
        self.user('admin@test.ci')
        admin = Client()
        self.login(admin, 'admin@test.ci')
        self.mail_ok = False
        page = admin.post('/admin/test-email', follow=True)
        self.assertContains(page, 'Clé refusée')
        self.assertContains(admin.get('/admin/agents'), 'Dernier échec d’envoi')
        self.mail_ok = True
        self.assertContains(admin.post('/admin/test-email', follow=True), 'E-mail de test envoyé')
        self.assertEqual(self.sent[-1]['to'], ['admin@test.ci'])

    def test_sans_service_d_envoi_la_page_comptes_previent_l_administrateur(self):
        self.user('admin@test.ci')
        admin = Client()
        self.login(admin, 'admin@test.ci')
        with override_settings(AGENTMAIL_API_KEY='', AGENTMAIL_INBOX=''):
            from ..services.mailer import reset_mailer
            reset_mailer()
            self.assertContains(admin.get('/admin/agents'), 'Aucun service d’envoi')

    def test_l_interface_d_administration_django_est_reservee_a_l_administrateur(self):
        self.user('awa@test.ci')
        self.user('admin@test.ci')
        employee, admin = Client(), Client()
        self.login(employee, 'awa@test.ci')
        self.login(admin, 'admin@test.ci')
        self.assertEqual(employee.get('/django-admin/').status_code, 302)
        self.assertEqual(admin.get('/django-admin/').status_code, 200)


class EmailSignupTests(BaseTest):
    values = {'nom': 'Kone', 'prenom': 'Issa', 'filiale': 'SNEDAI', 'email': 'issa@test.ci', 'password': PASSWORD, 'password_confirmation': PASSWORD, 'methode': 'email'}

    def test_inscription_par_code_a_6_chiffres(self):
        client = Client()
        response = client.post('/inscription', self.values)
        self.assertEqual(response['Location'], '/inscription/code')
        self.assertEqual(User.objects.count(), 0)
        self.assertIn('est votre code de confirmation', self.sent[-1]['subject'])
        self.assertEqual(self.sent[-1]['to'], ['issa@test.ci'])

        wrong = client.post('/inscription/code', {'code': '000000' if self.last_code() != '000000' else '111111'})
        self.assertContains(wrong, 'Il vous reste 4 essai', status_code=400)
        self.assertContains(client.post('/inscription/code/renvoyer'), 'Patientez encore', status_code=429)

        ok = client.post('/inscription/code', {'code': self.last_code()})
        self.assertEqual(ok['Location'], '/')
        user = User.objects.get()
        self.assertEqual((user.email, user.filiale), ('issa@test.ci', 'SNEDAI'))
        self.assertTrue(user.check_password(PASSWORD))
        self.assertEqual(client.get('/').status_code, 200)
        import time
        for _ in range(50):
            if any('Bienvenue' in m['subject'] for m in self.sent):
                break
            time.sleep(0.05)
        self.assertTrue(any('Bienvenue' in m['subject'] for m in self.sent))

    def test_cinq_codes_faux_bloquent_jusqu_a_un_nouveau_code(self):
        client = Client()
        client.post('/inscription', self.values)
        real = self.last_code()
        wrong = '999999' if real != '999999' else '888888'
        for _ in range(5):
            client.post('/inscription/code', {'code': wrong})
        self.assertContains(client.post('/inscription/code', {'code': real}), 'Trop d’essais', status_code=429)
        self.assertEqual(User.objects.count(), 0)

    def test_code_expire(self):
        client = Client()
        client.post('/inscription', self.values)
        session = client.session
        session['email_signup']['expires'] = 0
        session.save()
        self.assertContains(client.post('/inscription/code', {'code': self.last_code()}), 'expiré', status_code=400)

    def test_sans_service_d_envoi_l_inscription_par_e_mail_est_refusee(self):
        with override_settings(AGENTMAIL_API_KEY='', AGENTMAIL_INBOX=''):
            from ..services.mailer import reset_mailer
            reset_mailer()
            self.assertContains(Client().post('/inscription', self.values), 'pas encore disponible', status_code=400)

    def test_panne_d_envoi_du_code(self):
        self.mail_ok = False
        response = Client().post('/inscription', self.values)
        self.assertContains(response, 'n’a pas pu être envoyé', status_code=502)
        self.assertEqual(User.objects.count(), 0)

    def test_trop_de_tentatives_de_connexion(self):
        with override_settings(LOGIN_RATE_LIMIT=3):
            client = Client()
            codes = [client.post('/connexion', {'email': 'a@test.ci', 'password': 'x'}).status_code for _ in range(5)]
        self.assertEqual(codes, [401, 401, 401, 429, 429])


@override_settings(GOOGLE_CLIENT_ID='id', GOOGLE_CLIENT_SECRET='secret', GOOGLE_ENABLED=True)
class GoogleTests(BaseTest):
    def callback(self, client, profile):
        client.get('/auth/google')
        state = client.session['google_state']
        with mock.patch('evenements.views.auth.fetch_google_profile', return_value=profile):
            return client.get('/auth/google/callback', {'code': 'abc', 'state': state})

    def test_redirection_vers_google(self):
        response = Client().get('/auth/google')
        self.assertTrue(response['Location'].startswith('https://accounts.google.com/o/oauth2/v2/auth?'))
        self.assertIn('client_id=id', response['Location'])

    def test_creation_et_connexion_avec_google(self):
        client = Client()
        form = {'nom': 'Kone', 'prenom': 'Issa', 'filiale': 'SNEDAI', 'email': 'issa@test.ci', 'password': PASSWORD, 'password_confirmation': PASSWORD, 'methode': 'google'}
        self.assertEqual(client.post('/inscription', form)['Location'], '/auth/google')
        response = self.callback(client, {'id': 'g-1', 'email': 'issa@test.ci', 'verified': True})
        self.assertEqual(response['Location'], '/')
        user = User.objects.get()
        self.assertEqual(user.google_id, 'g-1')
        # reconnexion avec Google
        again = Client()
        self.assertEqual(self.callback(again, {'id': 'g-1', 'email': 'issa@test.ci', 'verified': True})['Location'], '/')

    def test_google_avec_une_autre_adresse_est_refuse(self):
        client = Client()
        form = {'nom': 'Kone', 'prenom': 'Issa', 'filiale': 'SNEDAI', 'email': 'issa@test.ci', 'password': PASSWORD, 'password_confirmation': PASSWORD, 'methode': 'google'}
        client.post('/inscription', form)
        response = self.callback(client, {'id': 'g-2', 'email': 'autre@test.ci', 'verified': True})
        self.assertContains(response, 'ne correspond pas', status_code=400)
        self.assertEqual(User.objects.count(), 0)

    def test_connexion_google_sans_compte_et_etat_invalide(self):
        self.assertContains(self.callback(Client(), {'id': 'g-3', 'email': 'x@test.ci', 'verified': True}), 'Créez d’abord', status_code=400)
        client = Client()
        client.get('/auth/google')
        self.assertContains(client.get('/auth/google/callback', {'code': 'abc', 'state': 'faux'}), 'annulée', status_code=400)

    def test_un_compte_existant_est_relie_a_google(self):
        self.user('awa@test.ci')
        self.callback(Client(), {'id': 'g-9', 'email': 'awa@test.ci', 'verified': True})
        self.assertEqual(User.objects.get().google_id, 'g-9')
